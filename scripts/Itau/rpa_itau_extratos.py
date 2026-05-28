#%%

from sqlalchemy import create_engine, text
import pandas as pd
from pathlib import Path
import configparser
from datetime import datetime
from itertools import chain
from openpyxl import load_workbook
from io import StringIO
from psycopg2 import sql
import re
import warnings

warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style, apply openpyxl's default"
)

DW_CONFIG_PATH = Path(r"E:\BI\config\config_datalake.ini")

dw = configparser.ConfigParser()
dw.read(DW_CONFIG_PATH, encoding="utf-8")

PG_HOST = dw.get("auth", "host", fallback=None)
PG_PORT = dw.get("auth", "port", fallback=None)
PG_DB = dw.get("auth", "db", fallback=None)
PG_USER = dw.get("auth", "user", fallback=None)
PG_PASS = dw.get("auth", "pwd", fallback=None)
SCHEMA = dw.get("auth", "schema", fallback="datalake")

caminho_base = Path(r"E:\RPA\RPA_Itau\downloads")
tabela = "itau_rpa_extratos"

engine = create_engine(
    f"postgresql+psycopg2://{PG_USER}:{PG_PASS}@{PG_HOST}:{PG_PORT}/{PG_DB}",
    pool_pre_ping=True
)

def listar_arquivos_itau(caminho_base: Path):

    padroes = [
        "extrato_PERIODO_*.xlsx",
        "extrato_periodo_*.xlsx"
    ]

    arquivos = chain.from_iterable(
        caminho_base.rglob(padrao)
        for padrao in padroes
    )

    return sorted(
        arquivo for arquivo in arquivos
        if arquivo.is_file()
    )


def obter_colunas_tabela(engine, schema: str, tabela: str):

    query = text("""
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = :schema
          AND table_name = :tabela
        ORDER BY ordinal_position
    """)

    with engine.begin() as conn:
        df_cols = pd.read_sql(
            query,
            conn,
            params={
                "schema": schema,
                "tabela": tabela
            }
        )

    return df_cols["column_name"].tolist()


def preparar_buffer_copy(df: pd.DataFrame):

    buffer = StringIO()

    df.to_csv(
        buffer,
        index=False,
        header=True,
        sep="\t",
        na_rep="\\N"
    )

    buffer.seek(0)

    return buffer


def substituir_apenas_arquivos_itau_listados_no_dw(
    df: pd.DataFrame,
    engine,
    schema: str,
    tabela: str
):

    if df.empty:
        print("DataFrame vazio. Nenhuma exclusão ou carga será feita.")
        return

    if "arquivo_origem" not in df.columns:
        print("Coluna arquivo_origem não encontrada no DataFrame.")
        print("Nenhuma exclusão ou carga será feita.")
        return

    arquivos_para_substituir = (
        df["arquivo_origem"]
        .dropna()
        .astype(str)
        .drop_duplicates()
        .tolist()
    )

    if not arquivos_para_substituir:
        print("Nenhum arquivo_origem válido encontrado no DataFrame.")
        print("Nenhuma exclusão ou carga será feita.")
        return

    buffer = preparar_buffer_copy(df)
    colunas = list(df.columns)

    raw_conn = engine.raw_connection()

    try:
        with raw_conn.cursor() as cursor:

            delete_sql = sql.SQL("""
                DELETE FROM {}.{}
                WHERE "arquivo_origem" = ANY(%s)
            """).format(
                sql.Identifier(schema),
                sql.Identifier(tabela)
            )

            cursor.execute(delete_sql, (arquivos_para_substituir,))

            linhas_deletadas = cursor.rowcount

            print(f"Arquivos Itaú encontrados/processados para substituir: {len(arquivos_para_substituir)}")
            print(f"Linhas apagadas no DW somente desses arquivos: {linhas_deletadas}")

            copy_sql = sql.SQL("""
                COPY {}.{} ({})
                FROM STDIN
                WITH (
                    FORMAT CSV,
                    HEADER TRUE,
                    DELIMITER E'\t',
                    NULL '\\N'
                )
            """).format(
                sql.Identifier(schema),
                sql.Identifier(tabela),
                sql.SQL(", ").join(sql.Identifier(col) for col in colunas)
            )

            cursor.copy_expert(copy_sql, buffer)

        raw_conn.commit()

        print("DELETE seletivo + INSERT de ITAÚ concluídos com sucesso.")

    except Exception as e:
        raw_conn.rollback()
        print("Erro na carga. O DELETE e o INSERT foram desfeitos.")
        raise e

    finally:
        raw_conn.close()


def identificar_aba_extrato(caminho_arquivo: Path):

    xls = pd.ExcelFile(caminho_arquivo, engine="openpyxl")

    for aba in xls.sheet_names:
        if aba.strip().lower() == "extrato detalhado":
            return xls, aba

    for aba in xls.sheet_names:
        if "extrato" in aba.strip().lower():
            return xls, aba

    return xls, xls.sheet_names[0]


def encontrar_linha_cabecalho(
    caminho_arquivo: Path,
    sheet_name: str,
    max_linhas=30
):

    wb = load_workbook(
        filename=caminho_arquivo,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb[sheet_name]

        for idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=max_linhas, values_only=True),
            start=0
        ):
            valores = [
                "" if valor is None else str(valor).strip().lower()
                for valor in row
            ]

            conteudo_linha = " | ".join(valores)

            if "data" in conteudo_linha and "descrição" in conteudo_linha:
                return idx

            if "data" in conteudo_linha and "descricao" in conteudo_linha:
                return idx

    finally:
        wb.close()

    return 4


def normalizar_colunas(df: pd.DataFrame):

    df.columns = (
        df.columns
        .astype(str)
        .str.replace("\n", " ", regex=False)
        .str.replace("\t", " ", regex=False)
        .str.replace("  ", " ", regex=False)
        .str.strip()
    )

    df = df.rename(columns={
        "Valor": "Valor (R$)",
        "Saldo": "Saldo (R$)",
        "Descricao": "Descrição",
        "ID Transacao": "ID Transação",
        "CPF CNPJ": "CPF/CNPJ",
        "CPF / CNPJ": "CPF/CNPJ",
        "Agencia/Conta": "Agência/Conta",
        "Agencia / Conta": "Agência/Conta",
        "Agência / Conta": "Agência/Conta"
    })

    return df


def extrair_metadados_nome_arquivo(nome_arquivo: str):

    padrao = re.compile(
        r"^extrato_PERIODO_\d+_\d{8}(?:_a)?_\d{8}_(?P<marca>.+)_(?P<data_arquivo>\d{8})_\d{6}\.xlsx$",
        re.IGNORECASE
    )

    match = padrao.match(nome_arquivo)

    if not match:
        return None

    marca = match.group("marca").strip("_")
    data_arquivo = match.group("data_arquivo")

    return {
        "marca": marca,
        "data_arquivo": data_arquivo
    }


# ============================================================
# CRIA ÍNDICE
# ============================================================

with engine.begin() as conn:
    conn.execute(text(f'''
        CREATE INDEX IF NOT EXISTS idx_{tabela}_arquivo_origem
        ON "{SCHEMA}"."{tabela}" ("arquivo_origem")
    '''))


# ============================================================
# LISTA TODOS OS ARQUIVOS ITAÚ
# ============================================================

arquivos_encontrados = listar_arquivos_itau(caminho_base)

print(f"Arquivos Itaú encontrados na pasta: {len(arquivos_encontrados)}")

for arquivo in arquivos_encontrados:
    print(f" - {arquivo.name}")


# ============================================================
# PROCESSA TODOS OS ARQUIVOS ENCONTRADOS
# ============================================================

dfs = []
data_carga = datetime.now()

colunas_desejadas = [
    "Data",
    "Descrição",
    "Valor (R$)",
    "Saldo (R$)",
    "ID Transação",
    "Nome",
    "CPF/CNPJ",
    "Instituição",
    "Agência/Conta"
]

for caminho_arquivo in arquivos_encontrados:
    arquivo = caminho_arquivo.name

    metadados = extrair_metadados_nome_arquivo(arquivo)

    if metadados is None:
        print(f"Fora do padrão esperado: {arquivo}")
        continue

    try:
        xls, nome_aba = identificar_aba_extrato(caminho_arquivo)

        print(f"Aba usada em {arquivo}: {nome_aba}")

        header_row = encontrar_linha_cabecalho(
            caminho_arquivo=caminho_arquivo,
            sheet_name=nome_aba
        )

        df = pd.read_excel(
            xls,
            sheet_name=nome_aba,
            header=header_row
        )

        df = normalizar_colunas(df)

        for coluna in colunas_desejadas:
            if coluna not in df.columns:
                df[coluna] = None

        df = df[colunas_desejadas]

        if "Descrição" not in df.columns:
            print(f"Coluna Descrição não encontrada: {arquivo}")
            continue

        df = df[df["Descrição"].notna()]
        df = df.dropna(how="all")

        if df.empty:
            print(f"Arquivo sem registros válidos: {arquivo}")
            continue

        df["marca"] = metadados["marca"]
        df["data_arquivo"] = metadados["data_arquivo"]
        df["arquivo_origem"] = arquivo
        df["data_atualizacao"] = data_carga

        dfs.append(df)

        print(f"✔ Processado: {arquivo} | Linhas: {len(df)}")

    except Exception as e:
        print(f"Erro ao processar {arquivo}: {e}")


# ============================================================
# CONSOLIDA E CARREGA NO DW
# ============================================================

if dfs:
    df_final = pd.concat(dfs, ignore_index=True)

    colunas_tabela = obter_colunas_tabela(
        engine=engine,
        schema=SCHEMA,
        tabela=tabela
    )

    for coluna in colunas_tabela:
        if coluna not in df_final.columns:
            df_final[coluna] = None

    colunas_validas = [
        coluna for coluna in colunas_tabela
        if coluna in df_final.columns
    ]

    colunas_ignoradas = [
        coluna for coluna in df_final.columns
        if coluna not in colunas_tabela
    ]

    if colunas_ignoradas:
        print("\nColunas ignoradas porque não existem na tabela:")
        print(colunas_ignoradas)

    df_final = df_final[colunas_validas]

    if df_final.empty:
        print("Após filtrar as colunas válidas, o DataFrame ficou vazio.")
        print("Nenhum DELETE ou INSERT foi executado no DW.")
    else:
        substituir_apenas_arquivos_itau_listados_no_dw(
            df=df_final,
            engine=engine,
            schema=SCHEMA,
            tabela=tabela
        )

        arquivos_processados = (
            df_final["arquivo_origem"]
            .dropna()
            .astype(str)
            .nunique()
        )

        print("\nCarga Itaú concluída.")
        print(f"Arquivos processados/substituídos: {arquivos_processados}")
        print(f"Linhas inseridas: {len(df_final)}")

else:
    df_final = pd.DataFrame()
    print("Nenhum arquivo válido encontrado para carregar.")


# %%