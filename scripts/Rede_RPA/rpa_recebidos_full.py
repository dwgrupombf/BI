#%%

"""
1) Reprocessa todos os arquivos *_RECEBIDOS_*.xlsx
2) Carrega na tabela datalake.rede_rpa_recebidos

Pasta: E:\RPA\RPA_Rede_2_0\downloads

"""

from sqlalchemy import create_engine, text
import pandas as pd
from pathlib import Path
import configparser
from itertools import chain
from datetime import datetime
from openpyxl import load_workbook
from io import StringIO
from psycopg2 import sql
import warnings

warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style, apply openpyxl's default"
)

DW_CONFIG_PATH = Path(r"E:\BI\config\config_datalake.ini")

dw = configparser.ConfigParser()
dw.read(DW_CONFIG_PATH, encoding="utf-8")

PG_HOST = dw.get("auth", "host", fallback=None)
PG_PORT = dw.get("auth", "port", fallback=None)
PG_DB = dw.get("auth", "db", fallback=None)
PG_USER = dw.get("auth", "user", fallback=None)
PG_PASS = dw.get("auth", "pwd", fallback=None)
SCHEMA = dw.get("auth", "schema", fallback="datalake")

caminho_base = Path(r"E:\RPA\RPA_Rede_2_0\downloads")
tabela = "rede_rpa_recebidos"
sheet_name = "pagamentos"

engine = create_engine(
    f"postgresql+psycopg2://{PG_USER}:{PG_PASS}@{PG_HOST}:{PG_PORT}/{PG_DB}",
    pool_pre_ping=True
)


def encontrar_linha_cabecalho(
    caminho_arquivo: Path,
    texto="data do recebimento",
    sheet_name="pagamentos",
    max_linhas=80
):

    if caminho_arquivo.suffix.lower() in [".xlsx", ".xlsm"]:
        wb = load_workbook(
            filename=caminho_arquivo,
            read_only=True,
            data_only=True
        )

        try:
            if sheet_name not in wb.sheetnames:
                print(f"Aba '{sheet_name}' não encontrada: {caminho_arquivo.name}")
                return None

            ws = wb[sheet_name]

            for idx, row in enumerate(
                ws.iter_rows(min_row=1, max_row=max_linhas, values_only=True),
                start=0
            ):
                conteudo_linha = " ".join(
                    "" if valor is None else str(valor).strip().lower()
                    for valor in row
                )

                if texto in conteudo_linha:
                    return idx

        finally:
            wb.close()

        return None

    else:
        df_amostra = pd.read_excel(
            caminho_arquivo,
            header=None,
            sheet_name=sheet_name,
            nrows=max_linhas,
            engine="xlrd"
        )

        for idx, row in df_amostra.iterrows():
            if texto in str(row.values).lower():
                return idx

        return None


def listar_arquivos_recebidos(caminho_base: Path):

    padroes = [
        "*_RECEBIDOS_*.xlsx",
        "*_RECEBIDOS_*.xlsm",
        "*_RECEBIDOS_*.xls"
    ]

    arquivos = chain.from_iterable(
        caminho_base.rglob(padrao)
        for padrao in padroes
    )

    return sorted(
        arquivo for arquivo in arquivos
        if arquivo.is_file()
    )


def obter_colunas_tabela(engine, schema: str, tabela: str):

    query = text("""
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = :schema
          AND table_name = :tabela
        ORDER BY ordinal_position
    """)

    with engine.begin() as conn:
        df_cols = pd.read_sql(
            query,
            conn,
            params={
                "schema": schema,
                "tabela": tabela
            }
        )

    return df_cols["column_name"].tolist()


def preparar_buffer_copy(df: pd.DataFrame):

    buffer = StringIO()

    df.to_csv(
        buffer,
        index=False,
        header=True,
        sep="\t",
        na_rep="\\N"
    )

    buffer.seek(0)

    return buffer


def normalizar_colunas_data(df: pd.DataFrame, colunas_data: list[str], arquivo: str):

    valores_sem_data = ["-", "", " ", "nan", "NaN", "None", "none"]

    for coluna in colunas_data:
        if coluna not in df.columns:
            continue

        serie_original = df[coluna]
        serie_limpa = serie_original.replace(valores_sem_data, pd.NA)
        serie_convertida = pd.to_datetime(
            serie_limpa,
            errors="coerce",
            dayfirst=True
        )

        valores_invalidos = (
            serie_limpa.notna()
            & serie_convertida.isna()
        )

        if valores_invalidos.any():
            qtd_invalidos = int(valores_invalidos.sum())
            exemplos = (
                serie_original[valores_invalidos]
                .astype(str)
                .drop_duplicates()
                .head(5)
                .tolist()
            )

            print(
                f"Aviso: {arquivo} | coluna '{coluna}' | "
                f"{qtd_invalidos} valor(es) invalido(s) convertido(s) para NULL: {exemplos}"
            )

        df[coluna] = serie_convertida

    return df


def normalizar_colunas_numericas(df: pd.DataFrame, colunas_numericas: list[str], arquivo: str):

    valores_sem_numero = ["-", "", " ", "nan", "NaN", "None", "none"]

    for coluna in colunas_numericas:
        if coluna not in df.columns:
            continue

        serie_original = df[coluna]
        serie_limpa = serie_original.replace(valores_sem_numero, pd.NA)
        serie_convertida = pd.to_numeric(serie_limpa, errors="coerce")

        valores_invalidos = (
            serie_limpa.notna()
            & serie_convertida.isna()
        )

        if valores_invalidos.any():
            qtd_invalidos = int(valores_invalidos.sum())
            exemplos = (
                serie_original[valores_invalidos]
                .astype(str)
                .drop_duplicates()
                .head(5)
                .tolist()
            )

            print(
                f"Aviso: {arquivo} | coluna '{coluna}' | "
                f"{qtd_invalidos} valor(es) invalido(s) convertido(s) para NULL: {exemplos}"
            )

        df[coluna] = serie_convertida

    return df


def normalizar_colunas_inteiras(df: pd.DataFrame, colunas_inteiras: list[str], arquivo: str):

    valores_sem_numero = ["-", "", " ", "nan", "NaN", "None", "none"]

    for coluna in colunas_inteiras:
        if coluna not in df.columns:
            continue

        serie_original = df[coluna]
        serie_limpa = serie_original.replace(valores_sem_numero, pd.NA)
        serie_numerica = pd.to_numeric(serie_limpa, errors="coerce")

        valores_invalidos_parse = (
            serie_limpa.notna()
            & serie_numerica.isna()
        )

        valores_nao_inteiros = (
            serie_numerica.notna()
            & (serie_numerica % 1 != 0)
        )

        valores_invalidos = valores_invalidos_parse | valores_nao_inteiros

        if valores_invalidos.any():
            qtd_invalidos = int(valores_invalidos.sum())
            exemplos = (
                serie_original[valores_invalidos]
                .astype(str)
                .drop_duplicates()
                .head(5)
                .tolist()
            )

            print(
                f"Aviso: {arquivo} | coluna '{coluna}' | "
                f"{qtd_invalidos} valor(es) invalido(s) para bigint convertido(s) para NULL: {exemplos}"
            )

        serie_numerica[valores_invalidos] = pd.NA
        df[coluna] = serie_numerica.astype("Int64")

    return df


def substituir_apenas_arquivos_listados_no_dw(
    df: pd.DataFrame,
    engine,
    schema: str,
    tabela: str
):

    if df.empty:
        print("DataFrame vazio. Nenhuma exclusão ou carga será feita.")
        return

    arquivos_para_substituir = (
        df["arquivo_origem"]
        .dropna()
        .astype(str)
        .drop_duplicates()
        .tolist()
    )

    if not arquivos_para_substituir:
        print("Nenhum arquivo_origem válido encontrado no DataFrame.")
        print("Nenhuma exclusão ou carga será feita.")
        return

    buffer = preparar_buffer_copy(df)
    colunas = list(df.columns)

    raw_conn = engine.raw_connection()

    try:
        with raw_conn.cursor() as cursor:

            delete_sql = sql.SQL("""
                DELETE FROM {}.{}
                WHERE "arquivo_origem" = ANY(%s)
            """).format(
                sql.Identifier(schema),
                sql.Identifier(tabela)
            )

            cursor.execute(delete_sql, (arquivos_para_substituir,))

            linhas_deletadas = cursor.rowcount

            print(f"Arquivos encontrados/processados para substituir: {len(arquivos_para_substituir)}")
            print(f"Linhas apagadas no DW somente desses arquivos: {linhas_deletadas}")

            copy_sql = sql.SQL("""
                COPY {}.{} ({})
                FROM STDIN
                WITH (
                    FORMAT CSV,
                    HEADER TRUE,
                    DELIMITER E'\t',
                    NULL '\\N'
                )
            """).format(
                sql.Identifier(schema),
                sql.Identifier(tabela),
                sql.SQL(", ").join(sql.Identifier(col) for col in colunas)
            )

            cursor.copy_expert(copy_sql, buffer)

        raw_conn.commit()

        print("DELETE seletivo + INSERT concluídos com sucesso.")

    except Exception as e:
        raw_conn.rollback()
        print("Erro na carga. O DELETE e o INSERT foram desfeitos.")
        raise e

    finally:
        raw_conn.close()


with engine.begin() as conn:
    conn.execute(text(f'''
        CREATE INDEX IF NOT EXISTS idx_{tabela}_arquivo_origem
        ON "{SCHEMA}"."{tabela}" ("arquivo_origem")
    '''))


arquivos_encontrados = listar_arquivos_recebidos(
    caminho_base=caminho_base
)

print(f"Arquivos RECEBIDOS encontrados na pasta: {len(arquivos_encontrados)}")

for arquivo in arquivos_encontrados:
    print(f" - {arquivo.parent.name}\\{arquivo.name}")


dfs = []
data_carga = datetime.now()

for caminho_arquivo in arquivos_encontrados:
    arquivo = caminho_arquivo.name
    pasta = caminho_arquivo.parent.name

    try:
        header_row = encontrar_linha_cabecalho(
            caminho_arquivo=caminho_arquivo,
            texto="data do recebimento",
            sheet_name=sheet_name
        )

        if header_row is None:
            print(f"Cabeçalho não encontrado: {arquivo}")
            continue

        engine_excel = (
            "openpyxl"
            if caminho_arquivo.suffix.lower() in [".xlsx", ".xlsm"]
            else "xlrd"
        )

        df = pd.read_excel(
            caminho_arquivo,
            skiprows=header_row,
            engine=engine_excel,
            sheet_name=sheet_name
        )

        df.columns = (
            df.columns
            .astype(str)
            .str.strip()
            .str.lower()
        )

        df = df.dropna(how="all")

        if df.empty:
            print(f"Arquivo sem linhas válidas: {arquivo}")
            continue

        df["marca"] = pasta
        df["arquivo_origem"] = arquivo
        df["data_atualizacao"] = data_carga

        df = normalizar_colunas_data(
            df=df,
            colunas_data=[
                "data do recebimento",
                "data original da venda",
                "data original de vencimento",
            ],
            arquivo=arquivo
        )
        df = normalizar_colunas_numericas(
            df=df,
            colunas_numericas=[
                "taxa mdr",
            ],
            arquivo=arquivo
        )
        df = normalizar_colunas_inteiras(
            df=df,
            colunas_inteiras=[
                "nsu/cv",
                "resumo de vendas/número do lote",
                "estabelecimento",
                "número de parcelas",
                "parcela",
            ],
            arquivo=arquivo
        )

        dfs.append(df)

        print(f"✔ Processado: {pasta}\\{arquivo} | Linhas: {len(df)}")

    except Exception as e:
        print(f"Erro ao processar {caminho_arquivo}: {e}")


if dfs:
    df_final = pd.concat(dfs, ignore_index=True)

    colunas_tabela = obter_colunas_tabela(
        engine=engine,
        schema=SCHEMA,
        tabela=tabela
    )

    colunas_validas = [
        coluna for coluna in df_final.columns
        if coluna in colunas_tabela
    ]

    colunas_ignoradas = [
        coluna for coluna in df_final.columns
        if coluna not in colunas_tabela
    ]

    if colunas_ignoradas:
        print("\nColunas ignoradas porque não existem na tabela:")
        print(colunas_ignoradas)

    df_final = df_final[colunas_validas]

    if df_final.empty:
        print("Após filtrar as colunas válidas, o DataFrame ficou vazio.")
        print("Nenhum DELETE ou INSERT foi executado no DW.")
    else:
        substituir_apenas_arquivos_listados_no_dw(
            df=df_final,
            engine=engine,
            schema=SCHEMA,
            tabela=tabela
        )

        arquivos_processados = (
            df_final["arquivo_origem"]
            .dropna()
            .astype(str)
            .nunique()
        )

        print(f"Arquivos processados/substituídos: {arquivos_processados}")
        print(f"Linhas inseridas: {len(df_final)}")

else:
    df_final = pd.DataFrame()
    print("Nenhum arquivo válido encontrado para carregar.")
    print("Nenhum DELETE ou INSERT foi executado no DW.")


# %%
