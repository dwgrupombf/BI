#%%

from sqlalchemy import create_engine, text, bindparam
import pandas as pd
from pathlib import Path
import configparser
from datetime import datetime
from openpyxl import load_workbook
from io import StringIO
from psycopg2 import sql


DW_CONFIG_PATH = Path(r"E:\BI\config\config_datalake.ini")
dw = configparser.ConfigParser()
dw.read(DW_CONFIG_PATH, encoding="utf-8")

PG_HOST = dw.get("auth", "host", fallback=None)
PG_PORT = dw.get("auth", "port", fallback=None)
PG_DB = dw.get("auth", "db", fallback=None)
PG_USER = dw.get("auth", "user", fallback=None)
PG_PASS = dw.get("auth", "pwd", fallback=None)
SCHEMA = dw.get("auth", "schema", fallback="datalake")

caminho_base = Path(r"E:\RPA\RPA_Rede_2_0\downloads")
tabela = "rede_rpa_a_receber"
sheet_name = "pagamentos futuros"

engine = create_engine(
    f"postgresql+psycopg2://{PG_USER}:{PG_PASS}@{PG_HOST}:{PG_PORT}/{PG_DB}",
    pool_pre_ping=True
)


def encontrar_linha_cabecalho(
    caminho_arquivo: Path,
    texto="data prevista do recebimento",
    sheet_name="pagamentos futuros",
    max_linhas=80
):

    if caminho_arquivo.suffix.lower() in [".xlsx", ".xlsm"]:
        wb = load_workbook(
            filename=caminho_arquivo,
            read_only=True,
            data_only=True
        )

        try:
            if sheet_name not in wb.sheetnames:
                print(f"Aba '{sheet_name}' não encontrada: {caminho_arquivo.name}")
                return None

            ws = wb[sheet_name]

            for idx, row in enumerate(
                ws.iter_rows(min_row=1, max_row=max_linhas, values_only=True),
                start=0
            ):
                conteudo_linha = " ".join(
                    "" if valor is None else str(valor).strip().lower()
                    for valor in row
                )

                if texto in conteudo_linha:
                    return idx

        finally:
            wb.close()

        return None

    else:

        df_amostra = pd.read_excel(
            caminho_arquivo,
            header=None,
            sheet_name=sheet_name,
            nrows=max_linhas,
            engine="xlrd"
        )

        for idx, row in df_amostra.iterrows():
            if texto in str(row.values).lower():
                return idx

        return None


def listar_arquivos_a_receber(caminho_base: Path):
    arquivos = []

    for caminho_arquivo in caminho_base.rglob("*"):
        if (
            caminho_arquivo.is_file()
            and "_A_RECEBER_" in caminho_arquivo.name
            and caminho_arquivo.suffix.lower() in [".xlsx", ".xlsm", ".xls"]
        ):
            arquivos.append(caminho_arquivo)

    return arquivos


def obter_colunas_tabela(engine, schema: str, tabela: str):
    query = text("""
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = :schema
          AND table_name = :tabela
        ORDER BY ordinal_position
    """)

    with engine.begin() as conn:
        df_cols = pd.read_sql(
            query,
            conn,
            params={
                "schema": schema,
                "tabela": tabela
            }
        )

    return df_cols["column_name"].tolist()


def copy_df_to_postgres(df: pd.DataFrame, engine, schema: str, tabela: str):

    if df.empty:
        print("DataFrame vazio. Nada para carregar.")
        return

    buffer = StringIO()

    df.to_csv(
        buffer,
        index=False,
        header=True,
        sep="\t",
        na_rep="\\N"
    )

    buffer.seek(0)

    colunas = list(df.columns)

    raw_conn = engine.raw_connection()

    try:
        with raw_conn.cursor() as cursor:
            copy_sql = sql.SQL("""
                COPY {}.{} ({})
                FROM STDIN
                WITH (
                    FORMAT CSV,
                    HEADER TRUE,
                    DELIMITER E'\t',
                    NULL '\\N'
                )
            """).format(
                sql.Identifier(schema),
                sql.Identifier(tabela),
                sql.SQL(", ").join(sql.Identifier(col) for col in colunas)
            )

            cursor.copy_expert(copy_sql, buffer)

        raw_conn.commit()

    except Exception as e:
        raw_conn.rollback()
        raise e

    finally:
        raw_conn.close()


with engine.begin() as conn:
    conn.execute(text(f'''
        CREATE INDEX IF NOT EXISTS idx_{tabela}_arquivo_origem
        ON "{SCHEMA}"."{tabela}" ("arquivo_origem")
    '''))


arquivos_encontrados = listar_arquivos_a_receber(caminho_base)
nomes_arquivos = [arquivo.name for arquivo in arquivos_encontrados]

print(f"Arquivos A_RECEBER encontrados na pasta: {len(nomes_arquivos)}")


if nomes_arquivos:
    query_arquivos = text(f'''
        SELECT DISTINCT "arquivo_origem"
        FROM "{SCHEMA}"."{tabela}"
        WHERE "arquivo_origem" IN :arquivos
    ''').bindparams(bindparam("arquivos", expanding=True))

    with engine.begin() as conn:
        arquivos_carregados = pd.read_sql(
            query_arquivos,
            conn,
            params={"arquivos": nomes_arquivos}
        )

    arquivos_carregados_set = set(
        arquivos_carregados["arquivo_origem"].astype(str)
    )

else:
    arquivos_carregados_set = set()


arquivos_novos = [
    arquivo for arquivo in arquivos_encontrados
    if arquivo.name not in arquivos_carregados_set
]

print(f"Arquivos novos para processar: {len(arquivos_novos)}")


dfs = []
data_carga = datetime.now()

for caminho_arquivo in arquivos_novos:
    arquivo = caminho_arquivo.name
    pasta = caminho_arquivo.parent.name

    try:
        header_row = encontrar_linha_cabecalho(
            caminho_arquivo=caminho_arquivo,
            texto="data prevista do recebimento",
            sheet_name=sheet_name
        )

        if header_row is None:
            print(f"Cabeçalho não encontrado: {arquivo}")
            continue

        engine_excel = (
            "openpyxl"
            if caminho_arquivo.suffix.lower() in [".xlsx", ".xlsm"]
            else "xlrd"
        )

        df = pd.read_excel(
            caminho_arquivo,
            skiprows=header_row,
            engine=engine_excel,
            sheet_name=sheet_name
        )

        df.columns = (
            df.columns
            .astype(str)
            .str.strip()
            .str.lower()
        )

        df = df.dropna(how="all")

        if df.empty:
            print(f"Arquivo sem linhas válidas: {arquivo}")
            continue

        df["marca"] = pasta
        df["arquivo_origem"] = arquivo
        df["data_atualizacao"] = data_carga

        if "data prevista do recebimento" in df.columns:
            df["data prevista do recebimento"] = pd.to_datetime(
                df["data prevista do recebimento"],
                errors="coerce"
            )

        dfs.append(df)

        print(f"✔ Processado: {pasta}\\{arquivo} | Linhas: {len(df)}")

    except Exception as e:
        print(f"Erro ao processar {caminho_arquivo}: {e}")


if dfs:
    df_final = pd.concat(dfs, ignore_index=True)

    colunas_tabela = obter_colunas_tabela(
        engine=engine,
        schema=SCHEMA,
        tabela=tabela
    )

    colunas_df = list(df_final.columns)

    colunas_validas = [
        coluna for coluna in colunas_df
        if coluna in colunas_tabela
    ]

    colunas_ignoradas = [
        coluna for coluna in colunas_df
        if coluna not in colunas_tabela
    ]

    if colunas_ignoradas:
        print("\nColunas ignoradas porque não existem na tabela:")
        print(colunas_ignoradas)

    df_final = df_final[colunas_validas]

    copy_df_to_postgres(
        df=df_final,
        engine=engine,
        schema=SCHEMA,
        tabela=tabela
    )

    print("\n🔥 Carga incremental concluída")
    print(f"Arquivos novos carregados: {len(dfs)}")
    print(f"Linhas inseridas: {len(df_final)}")

else:
    df_final = pd.DataFrame()
    print("Nenhum arquivo novo encontrado para carregar.")


# %%