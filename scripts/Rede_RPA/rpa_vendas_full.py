#%%

"""
1) Reprocessa todos os arquivos *Rede_Rel_Vendas*.xlsx
2) Carrega na tabela datalake.rede_rpa_vendas

Pasta: E:\RPA\RPA_Rede_2_0\downloads

"""

from sqlalchemy import create_engine, text
from itertools import chain
import pandas as pd
from pathlib import Path
import configparser
from datetime import datetime
import re
from openpyxl import load_workbook
from io import StringIO
from psycopg2 import sql
import warnings

warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style, apply openpyxl's default"
)

DW_CONFIG_PATH = Path(r"E:\BI\config\config_datalake.ini")

dw = configparser.ConfigParser()
dw.read(DW_CONFIG_PATH, encoding="utf-8")

PG_HOST = dw.get("auth", "host", fallback=None)
PG_PORT = dw.get("auth", "port", fallback=None)
PG_DB = dw.get("auth", "db", fallback=None)
PG_USER = dw.get("auth", "user", fallback=None)
PG_PASS = dw.get("auth", "pwd", fallback=None)
SCHEMA = dw.get("auth", "schema", fallback="datalake")

caminho_base = Path(r"E:\RPA\RPA_Rede_2_0\downloads")
tabela = "rede_rpa_vendas"

engine = create_engine(
    f"postgresql+psycopg2://{PG_USER}:{PG_PASS}@{PG_HOST}:{PG_PORT}/{PG_DB}",
    pool_pre_ping=True
)


def extrair_pv_nome_arquivo(nome_arquivo: str):

    match = re.search(r"^\d{2}_\d{4}_([0-9]+)_", nome_arquivo)
    return match.group(1) if match else None


def encontrar_linha_cabecalho(
    caminho_arquivo: Path,
    texto="data da venda",
    max_linhas=80
):

    wb = load_workbook(
        filename=caminho_arquivo,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb.active

        for idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=max_linhas, values_only=True),
            start=0
        ):
            conteudo_linha = " ".join(
                "" if valor is None else str(valor).strip().lower()
                for valor in row
            )

            if texto in conteudo_linha:
                return idx

    finally:
        wb.close()

    return None


def listar_arquivos_vendas(caminho_base: Path):

    padroes = [
        "*Rede_Rel_Vendas*.xlsx",
        "*Rede_Rel_Vendas*.xlsm"
    ]

    arquivos = chain.from_iterable(
        caminho_base.rglob(padrao)
        for padrao in padroes
    )

    return sorted(
        arquivo for arquivo in arquivos
        if arquivo.is_file()
    )


def obter_colunas_tabela(engine, schema: str, tabela: str):

    query = text("""
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = :schema
          AND table_name = :tabela
        ORDER BY ordinal_position
    """)

    with engine.begin() as conn:
        df_cols = pd.read_sql(
            query,
            conn,
            params={
                "schema": schema,
                "tabela": tabela
            }
        )

    return df_cols["column_name"].tolist()


def preparar_buffer_copy(df: pd.DataFrame):

    buffer = StringIO()

    df.to_csv(
        buffer,
        index=False,
        header=True,
        sep="\t",
        na_rep="\\N"
    )

    buffer.seek(0)

    return buffer


def substituir_apenas_arquivos_vendas_listados_no_dw(
    df: pd.DataFrame,
    engine,
    schema: str,
    tabela: str
):

    if df.empty:
        print("DataFrame vazio. Nenhuma exclusão ou carga será feita.")
        return

    if "arquivo_origem" not in df.columns:
        print("Coluna arquivo_origem não encontrada no DataFrame.")
        print("Nenhuma exclusão ou carga será feita.")
        return

    arquivos_para_substituir = (
        df["arquivo_origem"]
        .dropna()
        .astype(str)
        .drop_duplicates()
        .tolist()
    )

    if not arquivos_para_substituir:
        print("Nenhum arquivo_origem válido encontrado no DataFrame.")
        print("Nenhuma exclusão ou carga será feita.")
        return

    buffer = preparar_buffer_copy(df)
    colunas = list(df.columns)

    raw_conn = engine.raw_connection()

    try:
        with raw_conn.cursor() as cursor:

            delete_sql = sql.SQL("""
                DELETE FROM {}.{}
                WHERE "arquivo_origem" = ANY(%s)
            """).format(
                sql.Identifier(schema),
                sql.Identifier(tabela)
            )

            cursor.execute(delete_sql, (arquivos_para_substituir,))

            linhas_deletadas = cursor.rowcount

            print(f"Arquivos encontrados/processados para substituir: {len(arquivos_para_substituir)}")
            print(f"Linhas de VENDAS apagadas no DW somente desses arquivos: {linhas_deletadas}")

            copy_sql = sql.SQL("""
                COPY {}.{} ({})
                FROM STDIN
                WITH (
                    FORMAT CSV,
                    HEADER TRUE,
                    DELIMITER E'\t',
                    NULL '\\N'
                )
            """).format(
                sql.Identifier(schema),
                sql.Identifier(tabela),
                sql.SQL(", ").join(sql.Identifier(col) for col in colunas)
            )

            cursor.copy_expert(copy_sql, buffer)

        raw_conn.commit()

        print("DELETE seletivo + INSERT de VENDAS concluídos com sucesso.")

    except Exception as e:
        raw_conn.rollback()
        print("Erro na carga. O DELETE e o INSERT foram desfeitos.")
        raise e

    finally:
        raw_conn.close()


with engine.begin() as conn:
    conn.execute(text(f'''
        CREATE INDEX IF NOT EXISTS idx_{tabela}_arquivo_origem
        ON "{SCHEMA}"."{tabela}" ("arquivo_origem")
    '''))


arquivos_encontrados = listar_arquivos_vendas(
    caminho_base=caminho_base
)

print(f"Arquivos de VENDAS encontrados na pasta: {len(arquivos_encontrados)}")

for arquivo in arquivos_encontrados:
    print(f" - {arquivo.parent.name}\\{arquivo.name}")


dfs = []
data_carga = datetime.now()

for caminho_arquivo in arquivos_encontrados:
    arquivo = caminho_arquivo.name
    pasta = caminho_arquivo.parent.name

    try:
        header_row = encontrar_linha_cabecalho(
            caminho_arquivo=caminho_arquivo,
            texto="data da venda"
        )

        if header_row is None:
            print(f"Cabeçalho não encontrado: {arquivo}")
            continue

        df = pd.read_excel(
            caminho_arquivo,
            skiprows=header_row,
            engine="openpyxl"
        )

        df.columns = (
            df.columns
            .astype(str)
            .str.strip()
            .str.lower()
        )

        df = df.dropna(how="all")

        if df.empty:
            print(f"Arquivo sem linhas válidas: {arquivo}")
            continue

        df["marca"] = pasta
        df["arquivo_origem"] = arquivo
        df["pv"] = extrair_pv_nome_arquivo(arquivo)
        df["data_atualizacao"] = data_carga

        if "hora da venda" in df.columns:
            df["hora da venda"] = pd.to_datetime(
                df["hora da venda"],
                errors="coerce"
            ).dt.time

        dfs.append(df)

        print(
            f"✔ Processado: {pasta}\\{arquivo} | "
            f"PV: {df['pv'].iloc[0]} | "
            f"Linhas: {len(df)}"
        )

    except Exception as e:
        print(f"Erro ao processar {caminho_arquivo}: {e}")


if dfs:
    df_final = pd.concat(dfs, ignore_index=True)

    colunas_tabela = obter_colunas_tabela(
        engine=engine,
        schema=SCHEMA,
        tabela=tabela
    )

    colunas_validas = [
        coluna for coluna in df_final.columns
        if coluna in colunas_tabela
    ]

    colunas_ignoradas = [
        coluna for coluna in df_final.columns
        if coluna not in colunas_tabela
    ]

    if colunas_ignoradas:
        print("\nColunas ignoradas porque não existem na tabela:")
        print(colunas_ignoradas)

    df_final = df_final[colunas_validas]

    if df_final.empty:
        print("Após filtrar as colunas válidas, o DataFrame ficou vazio.")
        print("Nenhum DELETE ou INSERT foi executado no DW.")
    else:
        substituir_apenas_arquivos_vendas_listados_no_dw(
            df=df_final,
            engine=engine,
            schema=SCHEMA,
            tabela=tabela
        )

        arquivos_processados = (
            df_final["arquivo_origem"]
            .dropna()
            .astype(str)
            .nunique()
        )

        print(f"Arquivos processados/substituídos: {arquivos_processados}")
        print(f"Linhas inseridas: {len(df_final)}")

else:
    df_final = pd.DataFrame()
    print("Nenhum arquivo válido encontrado para carregar.")
    print("Nenhum DELETE ou INSERT foi executado no DW.")



# %%
