#%%

from sqlalchemy import create_engine, text
import pandas as pd
from pathlib import Path
import configparser
from datetime import datetime
import re
from openpyxl import load_workbook
from io import StringIO
from psycopg2 import sql

# ============================================================
# CONFIGURAÇÕES
# ============================================================

DW_CONFIG_PATH = Path(r"E:\BI\config\config_datalake.ini")

dw = configparser.ConfigParser()
dw.read(DW_CONFIG_PATH, encoding="utf-8")

PG_HOST = dw.get("auth", "host", fallback=None)
PG_PORT = dw.get("auth", "port", fallback=None)
PG_DB = dw.get("auth", "db", fallback=None)
PG_USER = dw.get("auth", "user", fallback=None)
PG_PASS = dw.get("auth", "pwd", fallback=None)
SCHEMA = dw.get("auth", "schema", fallback="datalake")

caminho_base = Path(r"E:\RPA\RPA_Rede_2_0\downloads")
tabela = "rede_rpa_vendas"

PERIODO_ARQUIVO = datetime.now().strftime("%m_%Y")
# PERIODO_ARQUIVO = "05_2026"

PADRAO_PERIODO = f"{PERIODO_ARQUIVO}_"

engine = create_engine(
    f"postgresql+psycopg2://{PG_USER}:{PG_PASS}@{PG_HOST}:{PG_PORT}/{PG_DB}",
    pool_pre_ping=True
)

# ============================================================
# FUNÇÕES
# ============================================================

def extrair_pv_nome_arquivo(nome_arquivo: str):

    match = re.search(r"^\d{2}_\d{4}_([0-9]+)_", nome_arquivo)
    return match.group(1) if match else None


def encontrar_linha_cabecalho(caminho_arquivo: Path, texto="data da venda", max_linhas=80):

    wb = load_workbook(
        filename=caminho_arquivo,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb.active

        for idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=max_linhas, values_only=True),
            start=0
        ):
            conteudo_linha = " ".join(
                "" if valor is None else str(valor).strip().lower()
                for valor in row
            )

            if texto in conteudo_linha:
                return idx

    finally:
        wb.close()

    return None


def listar_arquivos_vendas_periodo(caminho_base: Path, padrao_periodo: str):

    arquivos = []

    for caminho_arquivo in caminho_base.rglob("*"):
        if (
            caminho_arquivo.is_file()
            and caminho_arquivo.name.startswith(padrao_periodo)
            and "Rede_Rel_Vendas" in caminho_arquivo.name
            and caminho_arquivo.suffix.lower() in [".xlsx", ".xlsm"]
        ):
            arquivos.append(caminho_arquivo)

    return arquivos


def obter_colunas_tabela(engine, schema: str, tabela: str):
    query = text("""
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = :schema
          AND table_name = :tabela
        ORDER BY ordinal_position
    """)

    with engine.begin() as conn:
        df_cols = pd.read_sql(
            query,
            conn,
            params={
                "schema": schema,
                "tabela": tabela
            }
        )

    return df_cols["column_name"].tolist()


def preparar_buffer_copy(df: pd.DataFrame):
    buffer = StringIO()

    df.to_csv(
        buffer,
        index=False,
        header=True,
        sep="\t",
        na_rep="\\N"
    )

    buffer.seek(0)

    return buffer


def substituir_periodo_no_dw(
    df: pd.DataFrame,
    engine,
    schema: str,
    tabela: str,
    padrao_periodo: str
):

    if df.empty:
        print("DataFrame vazio. Nenhuma exclusão ou carga será feita.")
        return

    buffer = preparar_buffer_copy(df)
    colunas = list(df.columns)

    raw_conn = engine.raw_connection()

    try:
        with raw_conn.cursor() as cursor:

            delete_sql = sql.SQL("""
                DELETE FROM {}.{}
                WHERE "arquivo_origem" LIKE %s
            """).format(
                sql.Identifier(schema),
                sql.Identifier(tabela)
            )

            cursor.execute(delete_sql, [f"{padrao_periodo}%"])
            linhas_deletadas = cursor.rowcount

            print(f"Linhas apagadas no DW para {padrao_periodo}%: {linhas_deletadas}")

            copy_sql = sql.SQL("""
                COPY {}.{} ({})
                FROM STDIN
                WITH (
                    FORMAT CSV,
                    HEADER TRUE,
                    DELIMITER E'\t',
                    NULL '\\N'
                )
            """).format(
                sql.Identifier(schema),
                sql.Identifier(tabela),
                sql.SQL(", ").join(sql.Identifier(col) for col in colunas)
            )

            cursor.copy_expert(copy_sql, buffer)

        raw_conn.commit()

        print("DELETE + INSERT concluídos com sucesso.")

    except Exception as e:
        raw_conn.rollback()
        print("Erro na carga. O DELETE foi desfeito.")
        raise e

    finally:
        raw_conn.close()

with engine.begin() as conn:
    conn.execute(text(f'''
        CREATE INDEX IF NOT EXISTS idx_{tabela}_arquivo_origem
        ON "{SCHEMA}"."{tabela}" ("arquivo_origem")
    '''))

arquivos_encontrados = listar_arquivos_vendas_periodo(
    caminho_base=caminho_base,
    padrao_periodo=PADRAO_PERIODO
)

print(f"Período selecionado: {PERIODO_ARQUIVO}")
print(f"Arquivos encontrados para o período: {len(arquivos_encontrados)}")

for arquivo in arquivos_encontrados:
    print(f" - {arquivo.parent.name}\\{arquivo.name}")

dfs = []
data_carga = datetime.now()

for caminho_arquivo in arquivos_encontrados:
    arquivo = caminho_arquivo.name
    pasta = caminho_arquivo.parent.name

    try:
        header_row = encontrar_linha_cabecalho(
            caminho_arquivo=caminho_arquivo,
            texto="data da venda"
        )

        if header_row is None:
            print(f"Cabeçalho não encontrado: {arquivo}")
            continue

        df = pd.read_excel(
            caminho_arquivo,
            skiprows=header_row,
            engine="openpyxl"
        )

        df.columns = (
            df.columns
            .astype(str)
            .str.strip()
            .str.lower()
        )

        df = df.dropna(how="all")

        if df.empty:
            print(f"Arquivo sem linhas válidas: {arquivo}")
            continue

        df["marca"] = pasta
        df["arquivo_origem"] = arquivo
        df["pv"] = extrair_pv_nome_arquivo(arquivo)
        df["data_atualizacao"] = data_carga

        if "hora da venda" in df.columns:
            df["hora da venda"] = pd.to_datetime(
                df["hora da venda"],
                errors="coerce"
            ).dt.time

        dfs.append(df)

        print(
            f"✔ Processado: {pasta}\\{arquivo} | "
            f"PV: {df['pv'].iloc[0]} | "
            f"Linhas: {len(df)}"
        )

    except Exception as e:
        print(f"Erro ao processar {caminho_arquivo}: {e}")

if dfs:
    df_final = pd.concat(dfs, ignore_index=True)

    colunas_tabela = obter_colunas_tabela(
        engine=engine,
        schema=SCHEMA,
        tabela=tabela
    )

    colunas_validas = [
        coluna for coluna in df_final.columns
        if coluna in colunas_tabela
    ]

    colunas_ignoradas = [
        coluna for coluna in df_final.columns
        if coluna not in colunas_tabela
    ]

    if colunas_ignoradas:
        print("\nColunas ignoradas porque não existem na tabela:")
        print(colunas_ignoradas)

    df_final = df_final[colunas_validas]

    substituir_periodo_no_dw(
        df=df_final,
        engine=engine,
        schema=SCHEMA,
        tabela=tabela,
        padrao_periodo=PADRAO_PERIODO
    )

    print(f"Período: {PERIODO_ARQUIVO}")
    print(f"Arquivos processados: {len(dfs)}")
    print(f"Linhas inseridas: {len(df_final)}")

else:
    df_final = pd.DataFrame()
    print("Nenhum arquivo válido encontrado para carregar.")
    print("Nenhum DELETE foi executado no DW.")


# %%